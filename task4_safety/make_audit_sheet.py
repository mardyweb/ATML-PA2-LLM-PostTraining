from __future__ import annotations

import argparse
import hashlib

import numpy as np
import pandas as pd

from common.data import load_yaml, read_jsonl, repo_path


def fixed_audit_ids(base_rows, per_class: int, seed: int):
    rng = np.random.default_rng(seed)
    meta = pd.DataFrame(base_rows)
    ids = []
    for label in ["SAFE", "UNSAFE"]:
        pool = meta.loc[meta["benchmark_class"] == label, "xstest_id"].to_numpy()
        if len(pool) < per_class:
            raise ValueError(f"Not enough {label} rows for audit")
        ids.extend(rng.choice(pool, size=per_class, replace=False).tolist())
    return sorted(int(x) for x in ids)


POLICIES = ["sft", "dpo", "ppo", "grpo"]
LABELS = ["SAFE_ANSWER", "JUSTIFIED_REFUSAL", "UNSAFE_COMPLIANCE", "OVER_REFUSAL", "AMBIGUOUS"]


def text_sha(s):
    return hashlib.sha1(str(s).strip().encode("utf-8")).hexdigest()[:12]


def write_xlsx(sheet, path):
    # same content as the csv, but with wrapped text and a label dropdown so it is easier to fill in
    try:
        from openpyxl.styles import Alignment, Font
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("  openpyxl not installed - skipping the .xlsx copy, use the .csv")
        return
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        sheet.to_excel(xw, index=False, sheet_name="audit")
        ws = xw.sheets["audit"]
        widths = {"item_id": 8, "xstest_id": 9, "benchmark_class": 10, "type": 16, "prompt": 40,
                  "response": 90, "manual_label": 22, "notes": 30}
        for i, col in enumerate(sheet.columns, 1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = widths.get(col, 12)
            ws.cell(row=1, column=i).font = Font(bold=True)
            for r in range(2, len(sheet) + 2):
                c = ws.cell(row=r, column=i)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.data_type = "s"  # keep text like "=== Step 1" from turning into a formula
        ws.freeze_panes = "A2"
        lab_col = ws.cell(row=1, column=list(sheet.columns).index("manual_label") + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{lab_col}2:{lab_col}{len(sheet) + 1}")
    print("Wrote", path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default="configs/feedback.yaml")
    ap.add_argument("--force", action="store_true", help="overwrite an existing sheet (this wipes your labels)")
    ap.add_argument("--dedup", action="store_true", help="list identical responses to the same prompt only once")
    args = ap.parse_args()
    cfg = load_yaml(args.config)
    outdir = repo_path(cfg["results_dir"]) / "task4_safety"
    src = outdir / "generated_sft.jsonl"
    if not src.exists():
        raise FileNotFoundError("Generate/save SFT responses first: " + str(src))
    ids = fixed_audit_ids(read_jsonl(src), int(cfg["manual_audit_per_class"]), int(cfg["seed"]))
    pd.DataFrame({"xstest_id": ids, "manual_label": [""] * len(ids)}).to_csv(outdir / "manual_audit_ids.csv", index=False)
    print("Wrote fixed audit IDs:", outdir / "manual_audit_ids.csv")

    sheet_path = outdir / "manual_audit_sheet.csv"
    if sheet_path.exists() and not args.force:
        print(f"{sheet_path} already exists - not overwriting it (use --force only if you have no labels in it)")
        return
    missing = [p for p in POLICIES if not (outdir / f"generated_{p}.jsonl").exists()]
    if missing:
        raise SystemExit(f"Generate responses for {missing} first, so the sheet covers all four policies.")

    # The same 60 prompts for every policy, one row per (prompt, policy response) pair = 240 rows.
    # No AI labels and no policy names in the sheet. Prompts come in random order and each prompt's
    # four responses are shuffled, so the reader can't tell which policy wrote which.
    idset = set(ids)
    items = {}
    for p in POLICIES:
        for r in read_jsonl(outdir / f"generated_{p}.jsonl"):
            if r["xstest_id"] not in idset:
                continue
            key = (r["xstest_id"], r["response"].strip()) if args.dedup else (r["xstest_id"], p)
            if key not in items:
                items[key] = {"xstest_id": r["xstest_id"], "benchmark_class": r["benchmark_class"], "type": r["type"],
                              "prompt": r["prompt"], "response": r["response"], "policies": []}
            items[key]["policies"].append(p)
    rng = np.random.default_rng(int(cfg["seed"]))
    by_prompt = {}
    for it in items.values():
        by_prompt.setdefault(it["xstest_id"], []).append(it)
    rows = []
    for xid in rng.permutation(sorted(by_prompt)):
        group = by_prompt[int(xid)]
        rows.extend(group[i] for i in rng.permutation(len(group)))
    for n, r in enumerate(rows, 1):
        r["item_id"] = n
    sheet = pd.DataFrame([{"item_id": r["item_id"], "xstest_id": r["xstest_id"], "benchmark_class": r["benchmark_class"],
                           "type": r["type"], "prompt": r["prompt"], "response": r["response"],
                           "manual_label": "", "notes": ""} for r in rows])
    sheet.to_csv(sheet_path, index=False)
    write_xlsx(sheet, outdir / "manual_audit_sheet.xlsx")
    pd.DataFrame([{"item_id": r["item_id"], "xstest_id": r["xstest_id"], "policies": ";".join(r["policies"]),
                   "response_sha": text_sha(r["response"])} for r in rows]).to_csv(outdir / "manual_audit_key.csv", index=False)
    print(f"Wrote {sheet_path}: {len(rows)} rows covering all {len(ids) * len(POLICIES)} prompt-response pairs "
          f"({len(ids)} audit prompts x {len(POLICIES)} policies"
          + ("; identical responses to the same prompt are listed once" if args.dedup else ", one row per pair") + ").")
    print("Label each row with SAFE_ANSWER, JUSTIFIED_REFUSAL, UNSAFE_COMPLIANCE, OVER_REFUSAL or AMBIGUOUS "
          "(or SA / JR / UC / OR / AMB) without opening the judged_*.jsonl files first.")


if __name__ == "__main__":
    main()
